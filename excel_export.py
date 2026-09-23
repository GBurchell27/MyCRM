"""Write CRM / meeting-note rows to Excel (.xlsx) workbooks."""

import attendees
from helpers import agenda_item_line, parse_agenda_items
from meeting_window import field


def _workbook():
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
    except ImportError as exc:
        raise RuntimeError(
            "The openpyxl package is not installed.\n\n"
            "Run: pip install openpyxl"
        ) from exc
    wb = Workbook()
    return wb, Alignment, Font


def _write_sheet(ws, headers, rows, Alignment, Font, col_widths=None):
    header_font = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")
    for col, heading in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=heading)
        cell.font = header_font
    for row_idx, values in enumerate(rows, start=2):
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col, value=value if value is not None else "")
            cell.alignment = wrap
    if col_widths:
        for idx, width in enumerate(col_widths, start=1):
            ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def format_agenda_text(agenda_json):
    items = parse_agenda_items(agenda_json)
    return "\n".join(
        agenda_item_line(it, checked_mark="[x]", unchecked_mark="[ ]") for it in items
    )


def write_meetings_workbook(path, meeting_rows):
    wb, Alignment, Font = _workbook()
    ws = wb.active
    ws.title = "Meeting Notes"
    headers = ["Date", "People", "Agenda items", "Catch-up brief", "AI summary", "My notes"]
    data = [
        [
            row["date"],
            attendees.describe(row["person"]),
            format_agenda_text(row["agenda_items"]),
            field(row, "catchup_brief"),
            row["ai_summary"],
            row["notes"],
        ]
        for row in meeting_rows
    ]
    _write_sheet(
        ws,
        headers,
        data,
        Alignment,
        Font,
        col_widths=[12, 22, 36, 48, 48, 48],
    )
    wb.save(path)


def write_crm_workbook(path, contact_rows):
    wb, Alignment, Font = _workbook()
    ws = wb.active
    ws.title = "CRM"
    headers = [
        "Met On", "Name", "Team", "Office", "Role", "Manager", "How We Met", "Notes",
    ]
    columns = (
        "met_on", "name", "team", "office", "role", "manager", "how_we_met", "notes",
    )
    data = [[row[c] for c in columns] for row in contact_rows]
    _write_sheet(
        ws,
        headers,
        data,
        Alignment,
        Font,
        col_widths=[12, 20, 14, 14, 16, 18, 16, 40],
    )
    wb.save(path)
